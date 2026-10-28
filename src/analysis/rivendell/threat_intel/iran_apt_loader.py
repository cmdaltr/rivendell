"""The Iran APT threat intelligence workbook: XLSX sheets whose header row carries the ten columns
in threat_intel/schema.py (Evidential Element, Threat Group, ... Source Type).

Every sheet with that header is read, wherever the header row sits in the first rows (a title or
a note above it is fine). Columns may come in any order; extra columns are ignored, and a sheet
missing any of the ten is skipped with the reason, never half-read.
"""

import logging
from typing import List, Tuple

from rivendell.threat_intel.schema import COLUMNS, Indicator, from_row, unique

logger = logging.getLogger(__name__)

HEADER_SEARCH_ROWS = 20
DEFAULT_SOURCE = "Iran APT workbook"


class WorkbookError(ValueError):
    """The file isn't a workbook this loader can read."""


def _header(rows: List[Tuple]) -> Tuple[int, dict]:
    """(row index, {display name: column index}) for the first row holding every column."""
    wanted = {name.lower(): name for name in COLUMNS}
    for number, row in enumerate(rows):
        cells = {str(c).strip().lower(): i for i, c in enumerate(row) if c is not None}
        if all(name in cells for name in wanted):
            return number, {wanted[name]: cells[name] for name in wanted}
    return -1, {}


def load_workbook(path: str) -> List[Indicator]:
    """Every indicator in the workbook, deduplicated, with Source Type defaulting to the workbook."""
    try:
        import openpyxl
    except ImportError as exc:
        raise WorkbookError("reading XLSX needs openpyxl (requirements/base.txt)") from exc
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for a file that isn't XLSX
        raise WorkbookError(f"not a readable XLSX workbook: {exc}") from exc
    indicators = []
    try:
        for sheet in book.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            start, columns = _header(rows[:HEADER_SEARCH_ROWS])
            if start < 0:
                logger.info("Sheet '%s' has no row with all ten threat intelligence columns; skipped", sheet.title)
                continue
            for row in rows[start + 1:]:
                values = {name: (row[i] if i < len(row) else "") for name, i in columns.items()}
                indicator = from_row(values)
                if not indicator.source_type:
                    indicator = Indicator(**{**indicator.__dict__, "source_type": DEFAULT_SOURCE})
                indicators.append(indicator)
    finally:
        book.close()
    return unique(indicators)
